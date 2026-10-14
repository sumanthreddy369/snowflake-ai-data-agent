"""
Builds docs/workflow-and-tech-stack.pdf and .docx: the end-to-end workflow and
the full technology stack of snowflake-ai-data-agent, with where each piece
lives in the repo and how far each has been verified.

The content is written once, as renderer-neutral blocks (content()), and drawn
twice: as a PDF with reportlab (same fonts, colours and tables as
build_playbook.py) and as a Word document with python-docx. Needs reportlab,
python-docx and pypdfium2 (the diagram is rasterised for Word). Versions in the stack table are the ones installed and verified
in this repo's environment (and pinned by infra/terraform/.terraform.lock.hcl),
not "latest".
"""

import argparse
import logging
import re
import sys
import tempfile
from pathlib import Path

from reportlab.graphics.shapes import Drawing, Line, Polygon, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether as _RLKeepTogether
from reportlab.platypus import PageBreak as _RLPageBreak
from reportlab.platypus import SimpleDocTemplate
from reportlab.platypus import Spacer as _RLSpacer

# Same-folder helper module (docs/ isn't a package).
sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_playbook as pb  # noqa: E402
from build_playbook import ACCENT, INK, MUTED, register_fonts  # noqa: E402

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("build_workflow_doc")

DOCS_DIR = Path(__file__).resolve().parent
DEFAULT_STEM = DOCS_DIR / "workflow-and-tech-stack"

OURS = colors.HexColor("#EDE9FE")      # glue + guardrails we write
SNOW = colors.HexColor("#DBEAFE")      # inside Snowflake
EXT = colors.HexColor("#F3F4F6")       # outside Snowflake


def arch_diagram() -> Drawing:
    W, H = 7.5 * inch, 3.35 * inch
    d = Drawing(W, H)
    bw, bh, gap = 1.28 * inch, 0.5 * inch, 0.275 * inch
    x0 = (W - (5 * bw + 4 * gap)) / 2
    rows = [
        [("Alpaca API", "trades · bars · news", EXT), ("Kafka + DLQ", "streaming transport", EXT),
         ("Snowpipe", "Streaming + GCS batch", SNOW), ("Bronze", "raw, never rejects", SNOW),
         ("Silver", "Streams + Tasks", SNOW)],
        [("Gold", "dbt star schema", SNOW), ("Semantic View", "+ Cortex Search", SNOW),
         ("Cortex Agent", "analyst | search", SNOW), ("MCP · UI · API", "5 entry points", EXT),
         ("Analysts", "ask in English", EXT)],
    ]
    ys = [H - 0.62 * inch, H - 1.62 * inch]

    def arrow(x1, y1, x2, y2):
        d.add(Line(x1, y1, x2, y2, strokeColor=INK, strokeWidth=0.8))

    for r, row in enumerate(rows):
        y = ys[r]
        for i, (name, sub, fill) in enumerate(row):
            x = x0 + i * (bw + gap)
            d.add(Rect(x, y, bw, bh, rx=5, ry=5, fillColor=fill, strokeColor=ACCENT, strokeWidth=0.7))
            d.add(String(x + bw / 2, y + bh - 15, name, fontName="Body-Bold", fontSize=8.5, fillColor=INK, textAnchor="middle"))
            d.add(String(x + bw / 2, y + 8, sub, fontName="Body", fontSize=6.5, fillColor=MUTED, textAnchor="middle"))
            if i < 4:
                ax = x + bw + 3
                arrow(ax, y + bh / 2, ax + gap - 6, y + bh / 2)
                d.add(Polygon([ax + gap - 6, y + bh / 2 + 3, ax + gap - 2, y + bh / 2, ax + gap - 6, y + bh / 2 - 3],
                              fillColor=INK, strokeColor=INK))
    xs, xg = x0 + 4 * (bw + gap) + bw / 2, x0 + bw / 2
    ymid = (ys[0] + ys[1] + bh) / 2
    arrow(xs, ys[0], xs, ymid)
    arrow(xs, ymid, xg, ymid)
    arrow(xg, ymid, xg, ys[1] + bh + 4)
    d.add(Polygon([xg - 3, ys[1] + bh + 5, xg, ys[1] + bh + 1, xg + 3, ys[1] + bh + 5], fillColor=INK, strokeColor=INK))

    bands = [
        (0.82, "Wired and watched by:  Dagster · Terraform · GitHub Actions · Snowflake alerts · "
               "OpenTelemetry → Prometheus + Grafana + Jaeger", SNOW),
        (0.42, "Planned:  3 glue agents on Claude via Vertex AI Model Garden, behind one guardrail gateway "
               "(risk tiers · approvals · budgets · audit)", OURS),
    ]
    for yb, text, fill in bands:
        y = yb * inch
        d.add(Rect(x0, y, 5 * bw + 4 * gap, 0.3 * inch, rx=4, ry=4, fillColor=fill, strokeColor=ACCENT, strokeWidth=0.6))
        d.add(String(x0 + 8, y + 7.5, text, fontName="Body", fontSize=7, fillColor=INK))
    d.add(String(x0, 0.12 * inch, "Grey = outside Snowflake · Blue = inside Snowflake / existing tools · Purple = our planned glue + guardrails",
                 fontName="Body-Italic", fontSize=6.8, fillColor=MUTED))
    return d


def on_page(canvas, doc):
    canvas.saveState()
    canvas.setFont("Body", 7.5)
    canvas.setFillColor(MUTED)
    canvas.drawString(0.5 * inch, 0.4 * inch, "snowflake-ai-data-agent · Workflow and tech stack")
    canvas.drawRightString(8.0 * inch, 0.4 * inch, f"Page {doc.page}")
    canvas.restoreState()


# ------------------------------------------------------------------ blocks
# Content is a list of tuples so both renderers read the same thing:
#   ("p", markup, style) · ("note", markup) · ("bullets", [markup])
#   ("table", header, rows, widths_inches, bold_first) · ("diagram",)
#   ("pagebreak",) · ("spacer",) · ("keep", [blocks])
# Markup is the small reportlab subset used here: <b>, <i>, &lt; &gt; &amp;.

def P(text, style="body"):
    return ("p", text, style)


def note(text):
    return ("note", text)


def bullets(items):
    return [("bullets", list(items))]


def grid(header, rows, widths, bold_first=True):
    return ("table", list(header), [list(r) for r in rows], list(widths), bold_first)


def PageBreak():
    return ("pagebreak",)


def Spacer(*_):
    return ("spacer",)


def KeepTogether(blocks):
    return ("keep", list(blocks))


def diagram():
    return ("diagram",)


def stage(title, rows):
    """One workflow stage: a heading and a Step / What happens / Tool / Repo table."""
    return [P(title, "h2"), grid(["Step", "What happens", "Tool", "Where in the repo"], rows, [0.45, 3.45, 1.35, 2.25])]


def content() -> list:
    st = []

    # ------------------------------------------------------------------ overview
    st += [P("Workflow and Tech Stack", "title"),
           P("snowflake-ai-data-agent · real-time U.S. equities data on Snowflake, served to analysts through a governed AI "
             "agent, wired together with existing tools plus our own glue and guardrails", "subtitle")]
    st.append(P("What the project is", "h1"))
    st.append(P(
        "Live trades, minute bars and financial news flow from the Alpaca API through Kafka into Snowflake, where they are "
        "cleaned (Bronze → Silver), modelled into a star schema with dbt (Gold), and exposed through a semantic layer to "
        "Snowflake's Cortex Agent. Analysts ask questions in plain English through an MCP server, a Streamlit chat UI or a "
        "FastAPI service. Every path runs as the same role under the same row access policy, so the agent can never show a "
        "real-time price to someone not licensed to see it — everyone else sees prices delayed by 15 minutes."))
    st.append(P(
        "The design principle is the one companies follow: <b>use existing tools for the heavy lifting</b> (Snowflake, dbt, "
        "Dagster, Terraform, GitHub Actions, Cortex Agent) and write only <b>glue</b> (what connects them) and "
        "<b>guardrails</b> (what keeps them and the AI safe). The goal is to remove repetitive human work while keeping human "
        "approval for anything that changes money, access or production."))
    st.append(note(
        "<b>Status (14 Oct 2026):</b> verified locally — sample-data generators, DLQ routing, ONNX anomaly model, "
        "<i>dbt parse</i>, Dagster graph validation, <i>terraform validate</i>, and the local Kafka + observability stack "
        "(metrics, alerts and dashboard checked with replayed data); CI passes on GitHub. A Snowflake trial "
        "account exists and is being connected (key-pair auth set up). <b>No SQL has run against live Snowflake yet</b>; the "
        "glue agents and guardrail gateway are planned, not built."))
    st.append(Spacer(1, 6))
    st.append(diagram())

    # ------------------------------------------------------------------ workflow
    st.append(PageBreak())
    st.append(P("The workflow, end to end", "h1"))
    st.append(P("Eleven stages in the order they happen, from creating the infrastructure to answering an analyst's question "
                "and responding when something breaks. Each row names the tool doing the work and the file that defines it."))

    st += stage("A. Provision the infrastructure (once)", [
        ("A1", "Create the GCS bucket (no public access, 90-day file lifecycle), Pub/Sub topic + subscription, and the bucket "
               "notification that publishes every new file under raw/.", "Terraform + Google provider", "infra/terraform/gcp.tf"),
        ("A2", "Create the Snowflake database, the agent's warehouse ANALYST_WH (XSMALL, 30 s timeout), its credit monitor, "
               "six functional roles rolled up to SYSADMIN, and the GCS storage + Pub/Sub notification integrations.",
         "Terraform + Snowflake provider", "infra/terraform/snowflake.tf"),
        ("A3", "Grant Snowflake's own GCP service accounts read access to the bucket and subscription (replaces manual console "
               "steps).", "Terraform", "infra/terraform/gcp.tf"),
        ("A4", "Human step: review the Terraform plan and approve it.", "GitHub Actions (production env)",
         ".github/workflows/deploy.yml"),
    ])
    st += stage("B. Deploy the Snowflake objects (every release)", [
        ("B1", "Bronze: RAW_TRADES, RAW_BARS, RAW_SYMBOLS — every column VARCHAR so a load never rejects a row.",
         "Dagster → SQL", "sql/02_bronze/bronze_tables.sql"),
        ("B2", "Silver tables + Streams + Tasks; news table + Cortex Search service.", "Dagster → SQL",
         "sql/03_silver, sql/04_documents"),
        ("B3", "Stage, file format and auto-ingest pipes (run after B1–B2: they load those tables).", "Dagster → SQL",
         "sql/01_ingest/snowpipe_setup.sql"),
        ("B4", "Roles, GOLD/STAGING schemas, DELAYED_DATA_POLICY, grants (incl. FUTURE grants); then guardrails: agent "
               "warehouse, credit monitor, append-only audit log.", "Dagster → SQL", "sql/07_governance, sql/09_guardrails"),
        ("B5", "dbt build: staging views, then DIM_SYMBOLS, DIM_DATE, FCT_TRADES, FCT_BARS; post-hooks attach the policy; tests "
               "must pass.", "Dagster → dbt", "dbt/"),
        ("B6", "Semantic view → Cortex Agent → monitoring alerts.", "Dagster → SQL",
         "semantic_layer/, sql/10_agent, sql/11_monitoring"),
    ])
    st.append(P("The order is not the sql/ folder numbering: Dagster derives it from declared dependencies "
                "(orchestration/assets.py), so it can't drift. One command runs it all: "
                "<i>dagster job execute -m orchestration.definitions -j deploy_job</i>.", "body"))

    st += stage("C. Live ingestion (every second, market hours)", [
        ("C1", "Websocket receives trades and minute bars; each message is validated against a Pydantic contract.",
         "Python + Alpaca", "streaming/alpaca_stream_producer.py, schemas.py"),
        ("C2", "Valid messages → Kafka topics market-trades / market-bars; malformed ones → market-data-dlq with the error.",
         "Kafka", "streaming/alpaca_stream_producer.py"),
        ("C3", "Kafka Connect streams rows into Bronze within seconds.", "Snowpipe Streaming",
         "streaming/snowflake_kafka_connector.json"),
        ("C4", "Offline alternative: replay the sample CSVs onto the same topics (no Alpaca account needed).",
         "Python", "streaming/replay_sample_data.py"),
        ("C5", "Both emit OpenTelemetry metrics: messages per topic, DLQ count, event-to-Kafka lag, time of last message.",
         "OpenTelemetry", "streaming/telemetry.py"),
    ])
    st += stage("D. Batch path (nightly, after the close)", [
        ("D1", "Pull the day's bars, symbols and news from Alpaca REST into CSV.", "Dagster → Python",
         "streaming/backfill_historical.py, backfill_news.py"),
        ("D2", "Upload to gs://&lt;bucket&gt;/raw/&lt;feed&gt;/; the bucket publishes an event to Pub/Sub.", "Dagster + GCS",
         "orchestration/assets.py (gcs_raw_files)"),
        ("D3", "Snowpipe reads the event and loads the file into Bronze automatically; fills any gaps the stream left.",
         "Snowpipe auto-ingest", "sql/01_ingest/snowpipe_setup.sql"),
    ])
    st += stage("E. Transform", [
        ("E1", "Every minute (trades, bars) / 15 min (symbols, news): a Task reads the Stream of new Bronze rows and MERGEs "
               "them into Silver — typed, deduplicated on business keys, symbols upper-cased.", "Snowflake Tasks",
         "sql/03_silver/streams_and_tasks.sql"),
        ("E2", "Every 5 minutes in market hours: dbt rebuilds Gold — returns, 30-bar rolling volatility, is_suspect flag "
               "for >10% one-minute moves; policy re-attached; tests run.", "Dagster schedule → dbt", "dbt/models/marts/"),
        ("E3", "Blocking checks after each rebuild: policy still attached; failure stops anything downstream.",
         "Dagster asset checks", "orchestration/checks.py"),
    ])
    st += stage("F. Serve analysts", [
        ("F1", "Semantic view defines metrics, dimensions and synonyms over Gold; verified queries anchor common questions.",
         "Snowflake Semantic View", "semantic_layer/"),
        ("F2", "Cortex Agent routes each question: numbers → Cortex Analyst (SQL on the semantic view); "
               "\"why did it move\" → Cortex Search over news, with citations.", "Snowflake Cortex Agent", "sql/10_agent/"),
        ("F3", "Five entry points, all as role ANALYST_AGENT: direct Analyst, Cortex Agent, MCP server (3 tools), "
               "Streamlit chat, FastAPI.", "Python clients", "agent/, ui/, api/"),
    ])
    st += stage("G. Govern", [
        ("G1", "DELAYED_DATA_POLICY: rows newer than 15 minutes visible only to REALTIME_DESK; compared in UTC with SYSDATE() "
               "so it doesn't depend on session timezone.", "Row access policy", "sql/07_governance/"),
        ("G2", "Agent role is read-only on Gold; its warehouse has a 30 s timeout and a monthly credit cap; every question is "
               "written to an append-only audit log.", "Roles, resource monitor", "sql/09_guardrails/"),
    ])
    st += stage("H. Monitor and respond", [
        ("H1", "Snowflake ALERTs page on-call even if Dagster is down: ingestion stalled, policy detached (SEV-1); task "
               "failed, Gold stale (SEV-2).", "Snowflake alerts + email", "sql/11_monitoring/alerts.sql"),
        ("H2", "Guardrail checks every 15 minutes: policy attached, freshness, Bronze→Silver reconciliation, task health.",
         "Dagster", "orchestration/checks.py"),
        ("H3", "Outside Snowflake: the OTel Collector receives app metrics and traces; Prometheus scrapes it and "
               "kafka-exporter and alerts on producer stalled, DLQ > 1%, consumer lag > 10k, p95 event lag > 30 s.",
         "OpenTelemetry, Prometheus", "infra/local/"),
        ("H4", "Grafana dashboard of the live path (localhost:3001); Jaeger for traces (localhost:16687).",
         "Grafana, Jaeger", "infra/local/grafana/"),
        ("H5", "Human step: on-call follows the runbook; for governance incidents, pause the agent first.",
         "Runbook", "docs/runbook.md"),
    ])
    st += stage("I. Change a component", [
        ("I1", "Branch, edit, push: CI compiles all Python, runs the generators, validates the problem catalog, parses dbt, "
               "loads the Dagster graph, and validates Terraform.", "GitHub Actions", ".github/workflows/ci.yml"),
        ("I2", "Human step: code review, then approve the production deploy (Terraform plan → apply → Dagster deploy_job).",
         "GitHub Actions", ".github/workflows/deploy.yml"),
    ])
    st += stage("J. Control cost and shut down", [
        ("J1", "Account-wide cap: resource monitor stops every warehouse at 50 credits; warehouses auto-suspend after 60 s.",
         "Resource monitor", "sql/99_teardown/stop_spend.sql (level 0)"),
        ("J2", "Pause: suspend tasks, alerts and pipes; drop Cortex Search (costs credits while it exists). Remove: drop the "
               "database and integrations; terraform destroy.", "SQL + Terraform", "sql/99_teardown/stop_spend.sql (levels 1–2)"),
    ])
    st += stage("K. Planned: glue agents behind a guardrail gateway", [
        ("K1", "Analyst glue: asks the Cortex Agent, checks the answer came from executed SQL, returns it with as-of time and "
               "citations, or escalates.", "Claude on Vertex AI", "docs/glue-and-guardrails.md"),
        ("K2", "Incident glue: on an alert, checks systems in order, names the first broken link with evidence, proposes the "
               "runbook step.", "Claude on Vertex AI + MCP", "docs/glue-and-guardrails.md"),
        ("K3", "Change glue: drafts the dbt model and tests, opens a PR with a stage report, deploys after approval.",
         "Claude on Vertex AI + dbt/GitHub MCP", "docs/modeling-agents-study.md"),
        ("K4", "Gateway on every tool call: read/draft automatic; reversible prod change = 1 approval; irreversible, access or "
               "money = approval + reviewed plan; secrets, weakening checks, widening access blocked; budgets; audit.",
         "Python (planned)", "docs/glue-and-guardrails.md §4"),
        ("K5", "Model access: Claude through Vertex AI Model Garden in our GCP project (billing, IAM and audit in one cloud). "
               "Off by default; enabling the model in Model Garden is a human step. Analysts' questions stay on Snowflake "
               "Cortex so governed data never leaves Snowflake.", "Vertex AI + Anthropic SDK",
         "glue/llm.py, infra/terraform/vertex.tf"),
    ])

    # ----------------------------------------------------------------- freshness
    st.append(KeepTogether([
        P("Freshness budget", "h1"),
        grid(["Hop", "Target", "Mechanism"], [
            ("Alpaca → Kafka", "seconds", "Websocket producer"),
            ("Kafka → Bronze", "seconds", "Snowpipe Streaming"),
            ("Bronze → Silver", "≤ 1 minute", "Task every minute"),
            ("Silver → Gold", "≤ 5 minutes", "dbt every 5 minutes in market hours"),
            ("Gold → agent and non-entitled users", "+15 minutes (licensing, on purpose)", "DELAYED_DATA_POLICY"),
            ("Visible to the agent", "≈ 15–21 minutes; alert at 30", "—"),
            ("Visible to REALTIME_DESK", "≈ 1–6 minutes", "—"),
            ("Symbols / history / news", "daily / nightly / ≤ 1 hour", "Batch on purpose: they don't need to be real-time"),
        ], [2.3, 2.2, 3.0], bold_first=True),
    ]))

    # ----------------------------------------------------------------- stack
    st.append(PageBreak())
    st.append(P("Technology stack", "h1"))
    st.append(P("Versions are the ones installed and verified in this project's environment (Terraform providers pinned in "
                "infra/terraform/.terraform.lock.hcl). Snowflake and Alpaca are managed services, so they have no client version."))

    st.append(P("Data platform (Snowflake features used)", "h2"))
    st.append(grid(["Feature", "Used for", "Where"], [
        ("Snowpipe Streaming", "Row-level streaming from Kafka into Bronze in seconds", "streaming/snowflake_kafka_connector.json"),
        ("Snowpipe (auto-ingest) + GCS storage & notification integrations", "Batch files from the bucket into Bronze", "sql/01_ingest"),
        ("Streams + Tasks", "Change capture and scheduled MERGE, Bronze → Silver", "sql/03_silver"),
        ("Dynamic Tables", "Alternative to Streams+Tasks, shown side by side (not wired in)", "sql/03_silver/dynamic_tables_alternative.sql"),
        ("Semantic Views", "Governed metrics and dimensions the agent queries", "semantic_layer/semantic_view.sql"),
        ("Cortex Analyst", "Natural language → SQL over the semantic model", "agent/cortex_client.py"),
        ("Cortex Search", "Retrieval over news articles, with citations", "sql/04_documents"),
        ("Cortex Agents", "Routes each question to Analyst or Search", "sql/10_agent"),
        ("Row access policy", "15-minute delayed view for non-entitled roles", "sql/07_governance"),
        ("Roles + FUTURE grants + managed hierarchy", "Least privilege; grants survive dbt rebuilds", "sql/07_governance, infra/terraform"),
        ("Resource monitors + warehouse timeouts", "Credit caps; agent warehouse isolation", "sql/09_guardrails, sql/99_teardown"),
        ("ALERTs + email notification integration", "Paging that doesn't depend on the orchestrator", "sql/11_monitoring"),
        ("Key-pair (JWT) authentication", "Tools connect without passwords", "~/.snowflake/connections.toml (local)"),
    ], [2.5, 2.9, 2.1]))

    st.append(P("Tools and libraries", "h2"))
    st.append(grid(["Layer", "Technology", "Version", "Used for", "Where"], [
        ("Language", "Python", "3.12.10", "Everything outside SQL", "—"),
        ("Contracts", "Pydantic", "2.12.5", "Validating every record and agent response", "streaming/schemas.py, agent/schemas.py"),
        ("Source", "Alpaca API (websocket + REST)", "service", "Trades, bars, news", "streaming/"),
        ("Streaming", "Apache Kafka (KRaft) + Kafka Connect", "4.1.0 (Docker)", "Transport + DLQ", "infra/local/, streaming/"),
        ("Streaming client", "kafka-python · websockets", "3.0.11 · 16.1.1", "Producer and websocket", "streaming/"),
        ("Resilience", "tenacity · httpx", "9.1.4 · 0.28.1", "Retries with backoff; HTTP", "streaming/, agent/"),
        ("Transform", "dbt-core · dbt-snowflake", "1.12.5 · 1.12.1", "Gold star schema + tests", "dbt/"),
        ("Orchestration", "Dagster · dagster-dbt · dagster-snowflake", "1.13.25 · 0.29.25 · 0.29.25",
         "Deploy order, schedules, checks", "orchestration/"),
        ("Infrastructure", "Terraform · snowflakedb/snowflake · hashicorp/google", "≥1.6 · 2.21.0 · 7.46.1",
         "GCP + Snowflake account objects", "infra/terraform/"),
        ("Cloud storage", "Google Cloud Storage + Pub/Sub · google-cloud-storage", "service · 3.16.0",
         "Batch landing zone + events", "infra/terraform, orchestration/"),
        ("Snowflake access", "snowflake-connector-python · Snowflake CLI", "4.7.5 · 3.28.0", "Queries; local admin", "agent/, orchestration/"),
        ("Agent serving", "fastmcp · streamlit · fastapi · uvicorn", "4.0.10 · 1.64.0 · 0.142.1 · 0.54.0",
         "MCP server, chat UI, HTTP API", "agent/, ui/, api/"),
        ("Tracing (LLM)", "Langfuse", "4.15.6", "Tracing Cortex agent calls", "agent/"),
        ("Instrumentation", "OpenTelemetry SDK · OTLP HTTP exporter · FastAPI instrumentation", "1.45.0 · 1.45.0 · 0.66b0",
         "Metrics + traces from producer, replay, API", "streaming/telemetry.py, api/main.py"),
        ("Telemetry pipeline", "OpenTelemetry Collector (contrib)", "0.135.0", "Receives OTLP; exports to Prometheus + Jaeger",
         "infra/local/otel-collector.yaml"),
        ("Metrics + alerts", "Prometheus · kafka-exporter", "3.6.0 · 1.9.0", "Scraping; 4 alert rules", "infra/local/"),
        ("Dashboards", "Grafana", "12.2.0", "Live-pipeline dashboard", "infra/local/grafana/"),
        ("Traces", "Jaeger", "2.10.0", "Trace UI", "infra/local/"),
        ("Local runtime", "Docker Compose", "—", "Kafka + observability on one command", "infra/local/docker-compose.yml"),
        ("ML", "scikit-learn · skl2onnx · onnx", "1.9.1 · 1.20.0 · 1.23.1", "IsolationForest anomaly model → ONNX",
         "snowpark/train_anomaly_model.py"),
        ("Data handling", "pandas", "3.0.6", "Profiling, model training", "snowpark/"),
        ("CI/CD", "GitHub Actions", "service", "CI every push; approval-gated deploy", ".github/workflows/"),
        ("Docs", "reportlab · python-docx · pypdfium2 · PyYAML", "5.0.1 · 1.2.0 · 5.13.0 · 6.0.2",
         "These PDF and Word documents; the catalog", "docs/, knowledge/"),
        ("Models (glue agents)", "Claude on Vertex AI Model Garden · Anthropic SDK [vertex]", "service · 1.13.0",
         "Glue agents behind the guardrail gateway (off by default)", "glue/llm.py, infra/terraform/vertex.tf"),
    ], [0.95, 1.85, 1.35, 1.75, 1.6]))

    # ----------------------------------------------------------------- roles
    st.append(KeepTogether([
        P("Who can see what", "h1"),
        grid(["Role", "Purpose", "Sees real-time prices?"], [
            ("LOADER", "Ingestion (streaming + batch)", "n/a (writes Bronze only)"),
            ("TRANSFORMER", "dbt; can attach the policy (APPLY)", "n/a (builds Gold)"),
            ("ANALYST_AGENT", "The AI agent and all five entry points", "No — 15 minutes delayed"),
            ("BI_READER", "Human dashboards", "No — 15 minutes delayed"),
            ("REALTIME_DESK", "Entitled users (e.g. a trading desk)", "Yes"),
            ("AUDIT_REVIEWER", "Reads the agent audit log only", "n/a"),
        ], [1.6, 3.4, 2.5]),
    ]))

    # ----------------------------------------------------------------- humans
    st.append(KeepTogether([
        P("Where humans stay in the loop", "h1"),
        P("Repetitive work is automated; these decisions stay human on purpose:"),
        *bullets([
            "Approve the Terraform plan and every production deploy.",
            "Review code changes (CI must be green first).",
            "Grant real-time entitlement (REALTIME_DESK).",
            "Approve new metric definitions in the semantic layer.",
            "Respond to SEV-1 pages — pausing the agent is the safe default.",
            "Planned gateway: approve reversible production actions; approve + review irreversible ones.",
        ]),
    ]))

    # ----------------------------------------------------------------- status + run
    st.append(PageBreak())
    st.append(P("What's verified, written, and planned", "h1"))
    st.append(grid(["Component", "Status"], [
        ("Sample data generators, DLQ routing, ONNX model training", "Verified locally (run against sample data)"),
        ("CI pipeline", "Passing on GitHub"),
        ("dbt project", "dbt parse passes; not run against Snowflake"),
        ("Dagster asset graph and schedules", "dagster definitions validate passes; not run against Snowflake"),
        ("Terraform", "terraform validate passes against real providers; not applied"),
        ("All SQL (Bronze → Agent, governance, alerts, stop-spend)", "Written; not yet run on the live trial"),
        ("Snowflake trial account + key-pair connection", "Account active; key registration in progress"),
        ("Local Kafka + observability stack (OTel, Prometheus, Grafana, Jaeger)",
         "Verified locally: replayed data reached Prometheus and matched Kafka offsets; alerts and dashboard load"),
        ("Vertex AI wiring", "Client builds and terraform validates; no model call made (needs Model Garden enablement + budget)"),
        ("Approval-gated deploy workflow", "Written; not run"),
        ("Glue agents + guardrail gateway", "Planned (docs/glue-and-guardrails.md)"),
        ("pytest suite", "Not built yet"),
    ], [3.8, 3.7]))

    st.append(P("How to run it", "h1"))
    st.append(grid(["Goal", "Command"], [
        ("Install", "pip install -r requirements.txt"),
        ("Start Kafka + observability", "cd infra/local && docker compose up -d"),
        ("Send telemetry", "set OTEL_EXPORTER_OTLP_ENDPOINT=http://localhost:14318, then run the producer, replay or API"),
        ("Watch it", "Grafana localhost:3001 · Prometheus localhost:9091 · Jaeger localhost:16687"),
        ("Generate sample data", "python streaming/generate_sample_data.py · python streaming/generate_sample_news.py"),
        ("Provision infrastructure", "cd infra/terraform && terraform apply"),
        ("Deploy everything in order", "dagster job execute -m orchestration.definitions -j deploy_job"),
        ("Run schedules + UI", "dagster dev -m orchestration.definitions  (localhost:3000)"),
        ("Build Gold by hand", "cd dbt && dbt build"),
        ("Ask the agent", "python agent/cortex_agent_client.py \"What's happening with AAPL today?\""),
        ("Chat UI / API", "cd ui && streamlit run streamlit_app.py · cd api && uvicorn main:app"),
        ("Check Vertex AI wiring (no cost)", "python -m glue.llm --check"),
        ("Cap / pause spend", "Run sql/99_teardown/stop_spend.sql (levels 0–1) in Snowsight"),
        ("Rebuild these PDFs", "python docs/build_workflow_doc.py · python docs/build_playbook.py"),
    ], [2.0, 5.5]))

    st.append(P("Further reading in the repo", "h2"))
    st += bullets([
        "README.md — full technical reference, repository map and feature status.",
        "docs/production-readiness.md — existing tools vs our glue and guardrails; human work removed.",
        "docs/glue-and-guardrails.md — plan of record for the glue agents and gateway.",
        "docs/runbook.md — on-call procedures, one entry per alert.",
        "docs/snowflake-e2e-agent-playbook.pdf + knowledge/problem_catalog.yaml — 158 real-world problems and who handles each.",
    ])

    return st


# ------------------------------------------------------------------ PDF renderer

def _pdf_flowables(blocks):
    out = []
    for b in blocks:
        kind = b[0]
        if kind == "p":
            out.append(pb.P(b[1], b[2]))
        elif kind == "note":
            out.append(pb.note(b[1]))
        elif kind == "bullets":
            out += pb.bullets(b[1])
        elif kind == "table":
            out.append(pb.grid(b[1], b[2], b[3], bold_first=b[4]))
        elif kind == "diagram":
            out.append(arch_diagram())
        elif kind == "pagebreak":
            out.append(_RLPageBreak())
        elif kind == "spacer":
            out.append(_RLSpacer(1, 6))
        elif kind == "keep":
            out.append(_RLKeepTogether(_pdf_flowables(b[1])))
    return out


def build_pdf(out: Path, blocks: list) -> None:
    doc = SimpleDocTemplate(str(out), pagesize=letter, leftMargin=0.5 * inch, rightMargin=0.5 * inch,
                            topMargin=0.55 * inch, bottomMargin=0.6 * inch,
                            title="Workflow and Tech Stack", author="sumanthreddy369", subject="snowflake-ai-data-agent")
    doc.build(_pdf_flowables(blocks), onFirstPage=on_page, onLaterPages=on_page)
    logger.info("wrote %s", out)


# ------------------------------------------------------------------ Word renderer

_TOKEN = re.compile(r"(</?b>|</?i>)")
_ENTITIES = {"&lt;": "<", "&gt;": ">", "&amp;": "&", "&nbsp;": "\u00a0"}


def _unescape(text: str) -> str:
    for k, v in _ENTITIES.items():
        text = text.replace(k, v)
    return text


def _add_markup(par, markup: str, size=None, bold=False, color=None) -> None:
    """Write reportlab-style inline markup into a python-docx paragraph as runs."""
    b, i = bold, False
    for part in _TOKEN.split(markup):
        if part in ("<b>", "</b>"):
            b = part == "<b>" or bold
            continue
        if part in ("<i>", "</i>"):
            i = part == "<i>"
            continue
        if not part:
            continue
        run = par.add_run(_unescape(part))
        run.bold, run.italic = b, i
        if size:
            run.font.size = size
        if color is not None:
            run.font.color.rgb = color


def _shade(cell, hex_fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _diagram_png(path: Path) -> None:
    from reportlab.graphics import renderPDF
    import pypdfium2 as pdfium
    pdf_path = path.with_suffix(".pdf")
    renderPDF.drawToFile(arch_diagram(), str(pdf_path))
    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        pdf[0].render(scale=3).to_pil().save(str(path))
    finally:
        pdf.close()  # Windows can't delete the temp dir while the file is open


def _docx_blocks(doc, blocks, tmp: Path) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Inches, Pt, RGBColor
    white = RGBColor(0xFF, 0xFF, 0xFF)
    for b in blocks:
        kind = b[0]
        if kind == "p":
            style = {"title": "Title", "h1": "Heading 1", "h2": "Heading 2"}.get(b[2])
            par = doc.add_paragraph(style=style) if style else doc.add_paragraph()
            if b[2] == "subtitle":
                _add_markup(par, b[1], size=Pt(12), color=RGBColor(0x6B, 0x72, 0x80))
            else:
                _add_markup(par, b[1])
        elif kind == "note":
            t = doc.add_table(rows=1, cols=1)
            cell = t.rows[0].cells[0]
            cell.width = Inches(7.3)
            _shade(cell, "EFF6FF")
            _add_markup(cell.paragraphs[0], b[1], size=Pt(9.5))
            doc.add_paragraph()
        elif kind == "bullets":
            for item in b[1]:
                _add_markup(doc.add_paragraph(style="List Bullet"), item)
        elif kind == "table":
            header, rows, widths, bold_first = b[1], b[2], b[3], b[4]
            scale = 7.3 / sum(widths)
            t = doc.add_table(rows=1 + len(rows), cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            for r, values in enumerate([header] + rows):
                for c, value in enumerate(values):
                    cell = t.rows[r].cells[c]
                    cell.width = Inches(widths[c] * scale)
                    par = cell.paragraphs[0]
                    if r == 0:
                        _shade(cell, "1E3A8A")
                        _add_markup(par, value, size=Pt(8.5), bold=True, color=white)
                    else:
                        if r % 2 == 0:
                            _shade(cell, "F3F4F6")
                        _add_markup(par, value, size=Pt(8.5), bold=(c == 0 and bold_first))
            doc.add_paragraph()
        elif kind == "diagram":
            png = tmp / "architecture.png"
            _diagram_png(png)
            doc.add_picture(str(png), width=Inches(7.3))
        elif kind == "pagebreak":
            doc.add_page_break()
        elif kind == "keep":
            _docx_blocks(doc, b[1], tmp)


def build_docx(out: Path, blocks: list) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.6))
    sec.top_margin, sec.bottom_margin = Inches(0.6), Inches(0.6)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name, st.font.size = "Arial", Pt(size)
        st.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) if name != "Title" else RGBColor(0x1F, 0x29, 0x37)

    footer = sec.footer.paragraphs[0]
    footer.text = "snowflake-ai-data-agent · Workflow and tech stack"
    footer.runs[0].font.size = Pt(8)

    doc.core_properties.title = "Workflow and Tech Stack"
    doc.core_properties.author = "sumanthreddy369"

    with tempfile.TemporaryDirectory() as tmp:
        _docx_blocks(doc, blocks, Path(tmp))
        doc.save(str(out))
    logger.info("wrote %s", out)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out-stem", type=Path, default=DEFAULT_STEM,
                        help="Output path without extension; writes <stem>.pdf and <stem>.docx")
    parser.add_argument("--format", choices=["pdf", "docx", "both"], default="both")
    args = parser.parse_args()
    blocks = content()
    if args.format in ("pdf", "both"):
        register_fonts()
        build_pdf(args.out_stem.with_suffix(".pdf"), blocks)
    if args.format in ("docx", "both"):
        build_docx(args.out_stem.with_suffix(".docx"), blocks)


if __name__ == "__main__":
    main()
